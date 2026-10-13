#!/usr/bin/env python3
"""Build the Sprint 1 Consultation deck for IS212 SPM (ConnectSphere, Team G6T4)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- palette ----
INK    = RGBColor(0x0F, 0x17, 0x2A)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
SLATE  = RGBColor(0x47, 0x55, 0x69)
MUTED  = RGBColor(0x94, 0xA3, 0xB8)
LIGHT  = RGBColor(0xF1, 0xF5, 0xF9)
CARD   = RGBColor(0xF8, 0xFA, 0xFC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x16, 0xA3, 0x4A)
AMBER  = RGBColor(0xD9, 0x77, 0x06)
FONT   = "Helvetica Neue"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def _set(run, size, color, bold=False, italic=False, font=FONT):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.italic = italic; run.font.name = font

def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    return s

def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line_spacing=None):
    """runs: list of paragraphs; each paragraph is list of (txt,size,color,bold,italic)."""
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after)
        if line_spacing: p.line_spacing = line_spacing
        for seg in para:
            txt, size, color = seg[0], seg[1], seg[2]
            bold = seg[3] if len(seg) > 3 else False
            italic = seg[4] if len(seg) > 4 else False
            r = p.add_run(); r.text = txt; _set(r, size, color, bold, italic)
    return tb

def content_slide(num, kicker, title):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, 0, 0, Inches(0.16), SH, INDIGO)          # left accent bar
    text(s, Inches(0.7), Inches(0.42), Inches(11), Inches(0.35),
         [[(f"{num:02d}  ·  {kicker.upper()}", 12, INDIGO, True)]])
    text(s, Inches(0.7), Inches(0.72), Inches(12), Inches(0.7),
         [[(title, 30, INK, True)]])
    rect(s, Inches(0.72), Inches(1.42), Inches(1.5), Inches(0.05), INDIGO)
    # footer
    text(s, Inches(0.7), Inches(7.02), Inches(9), Inches(0.35),
         [[("ConnectSphere EPVBS  ·  IS212 SPM  ·  Team G6T4", 9, MUTED)]])
    text(s, Inches(11.8), Inches(7.02), Inches(1.2), Inches(0.35),
         [[(str(num), 9, MUTED)]], align=PP_ALIGN.RIGHT)
    return s

def bullets(slide, x, y, w, h, items, size=15, gap=7):
    """items: list of (text, level) or (text, level, color)."""
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        txt, lvl = it[0], it[1]
        color = it[2] if len(it) > 2 else SLATE
        bold = it[3] if len(it) > 3 else False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        bullet = "•   " if lvl == 0 else "–   "
        indent = "" if lvl == 0 else "      "
        r = p.add_run(); r.text = indent + bullet + txt
        _set(r, size if lvl == 0 else size-1, color, bold)
    return tb

def placeholder(slide, x, y, w, h, label):
    box = rect(slide, x, y, w, h, CARD, line=MUTED)
    box.line.dash_style = None
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; _set(r, 12, MUTED, italic=True)
    return box

def metric_card(slide, x, y, w, h, number, label, ncolor=INDIGO):
    rect(slide, x, y, w, h, CARD, line=LIGHT)
    text(slide, x, y+Inches(0.22), w, Inches(0.7),
         [[(number, 34, ncolor, True)]], align=PP_ALIGN.CENTER)
    text(slide, x, y+h-Inches(0.72), w, Inches(0.6),
         [[(label, 11, SLATE)]], align=PP_ALIGN.CENTER)

# ============================================================ 1. TITLE
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, INK)
rect(s, 0, Inches(3.52), SW, Inches(0.06), INDIGO)
text(s, Inches(0.9), Inches(0.8), Inches(6), Inches(0.4),
     [[("IS212 · SOFTWARE PROJECT MANAGEMENT", 13, MUTED, True)]])
text(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.4),
     [[("ConnectSphere EPVBS", 52, WHITE, True)]])
text(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(0.8),
     [[("Sprint 1 — Scrum Process Consultation (Week 7)", 22, RGBColor(0xC7,0xD2,0xFE))]])
text(s, Inches(0.9), Inches(5.3), Inches(11.5), Inches(1.6),
     [[("Team G6T4  ·  [ member names + Scrum roles — PO / SM / Dev ]", 15, MUTED)],
      [("1 Oct 2026  ·  Section G6", 13, MUTED)]], space_after=8)

# ============================================================ 2. AGENDA
s = content_slide(2, "Agenda", "What we'll walk you through")
col1 = [
    ("Product backlog — structure & scope", 0, INK, True),
    ("How we estimate story points", 0, INK, True),
    ("C4 architecture diagrams", 0, INK, True),
    ("How the sprint was conducted", 0, INK, True),
]
col2 = [
    ("Blockers & how we resolved them", 0, INK, True),
    ("Test cases in the CI pipeline", 0, INK, True),
    ("Definition of Done — & where it lives", 0, INK, True),
    ("Unit testing — location, auto vs manual", 0, INK, True),
]
bullets(s, Inches(0.7), Inches(2.0), Inches(6.2), Inches(4.5), col1, size=18, gap=20)
bullets(s, Inches(7.0), Inches(2.0), Inches(6.0), Inches(4.5), col2, size=18, gap=20)
text(s, Inches(0.7), Inches(6.4), Inches(12), Inches(0.5),
     [[("Everything below is shown live — Jira, GitHub & CI — this deck is just the map.", 13, INDIGO, False, True)]])

# ============================================================ 3. AT A GLANCE
s = content_slide(3, "Sprint 1", "Sprint 1 at a glance")
y = Inches(2.0); w = Inches(2.78); h = Inches(1.9); gap = Inches(0.18); x0 = Inches(0.7)
metric_card(s, x0, y, w, h, "75", "backlog items")
metric_card(s, x0+(w+gap), y, w, h, "240", "story points (total)")
metric_card(s, x0+2*(w+gap), y, w, h, "16", "stories Done (~49 pts)")
metric_card(s, x0+3*(w+gap), y, w, h, "272", "automated tests in CI", GREEN)
bullets(s, Inches(0.7), Inches(4.3), Inches(12), Inches(2.4), [
    ("Scope — 12 epics mapped to the 20 core features; MoSCoW prioritised, DEEP / just-in-time refinement.", 0),
    ("Velocity — [ fill in: committed vs completed pts for Sprint 1 ]  ·  Sprint length: 2 weeks.", 0, AMBER),
    ("Quality — every merged story passed CI (lint + build + unit + E2E) and met the Definition of Done.", 0),
    ("Tech stack — Next.js (React) + Tailwind frontend, Python Flask REST API, Supabase (Postgres + Auth).", 0),
], size=15, gap=11)

# ============================================================ 4. PRODUCT BACKLOG
s = content_slide(4, "Item 1 · Backlog", "Product Backlog")
bullets(s, Inches(0.7), Inches(1.9), Inches(7.2), Inches(5), [
    ("Single product backlog for the whole project, derived from the Customer Brief.", 0),
    ("12 epics → user stories; scoped to the 20 core features (first release).", 0),
    ("Full traceability: every core feature maps to specific stories.", 0),
    ("MoSCoW — the 20 core features are Must; extras are Should / Could / Won't.", 0),
    ("e.g. pricing / payment = Won't this release.", 1),
    ("DEEP backlog — top items fully detailed with ACs; lower items stay coarse,", 0),
    ("refined just-in-time (avoids big-design-up-front).", 1),
    ("Living artefact — refined after each customer Q&A; conflicts J1 / J2 logged, not assumed.", 0, INK, True),
], size=14.5, gap=9)
placeholder(s, Inches(8.3), Inches(1.9), Inches(4.3), Inches(3.3),
            "[ screenshot / live: Jira backlog view\n(epics + stories) ]")
text(s, Inches(8.3), Inches(5.35), Inches(4.3), Inches(1.2),
     [[("LIVE: ", 12, INDIGO, True), ("Jira board — spm-g6t4", 12, SLATE)],
      [("DOC: ", 12, INDIGO, True), ("Product_Backlog.md (authoring + coverage table)", 11, SLATE)]],
     space_after=6)

# ============================================================ 5. ESTIMATION
s = content_slide(5, "Item 2 · Estimation", "How we estimate story points")
bullets(s, Inches(0.7), Inches(1.9), Inches(7.2), Inches(5), [
    ("Planning Poker — the whole team estimates together, simultaneous reveal.", 0),
    ("Modified Fibonacci scale: 1, 2, 3, 5, 8  (we cap at 8 → anything bigger is split).", 0),
    ("Relative sizing, not hours — points capture effort + complexity + uncertainty vs a reference story.", 0),
    ("Discuss on disagreement — outliers explain; re-vote until converged (shared understanding > the number).", 0),
    ("Estimates stored in Jira (Story Points field) — the single source of truth.", 0, INK, True),
    ("Only refined / sprint-ready stories are pointed — coarse backlog items stay unestimated.", 0),
], size=15, gap=11)
placeholder(s, Inches(8.3), Inches(1.9), Inches(4.3), Inches(2.6),
            "[ optional image: planning-poker\nFibonacci cards 1·2·3·5·8 ]")
text(s, Inches(8.3), Inches(4.65), Inches(4.3), Inches(1.8),
     [[("Reference anchors", 13, INDIGO, True)],
      [("8 pts — Venue availability calendar", 11, SLATE)],
      [("5 pts — Create an event request", 11, SLATE)],
      [("1 pt  — Submit a completed request", 11, SLATE)]], space_after=6)

# ============================================================ 6. C4
s = content_slide(6, "Item 3 · Architecture", "C4 Architecture Diagrams")
text(s, Inches(0.7), Inches(1.8), Inches(12), Inches(0.5),
     [[("We model C1 System Context + C2 Container (per Week 4), plus C4 code-level for a key story.", 14, SLATE)]])
pw = Inches(3.9); ph = Inches(2.9); py = Inches(2.5)
placeholder(s, Inches(0.7), py, pw, ph, "[ image: C1 — System Context ]")
placeholder(s, Inches(4.75), py, pw, ph, "[ image: C2 — Container ]")
placeholder(s, Inches(8.8), py, pw, ph, "[ image: C4 L4 — SCRUM-32\nclass + sequence ]")
text(s, Inches(0.7), py+ph+Inches(0.05), pw, Inches(0.8),
     [[("5 roles · ConnectSphere · Supabase + Email (external)", 11, SLATE)]])
text(s, Inches(4.75), py+ph+Inches(0.05), pw, Inches(0.8),
     [[("Web App (Next.js) → Backend API (Flask) → Supabase", 11, SLATE)]])
text(s, Inches(8.8), py+ph+Inches(0.05), pw, Inches(0.8),
     [[("Booking approval: classes + sequence (SVG, in repo)", 11, SLATE)]])
text(s, Inches(0.7), Inches(6.75), Inches(12), Inches(0.4),
     [[("SOURCE: ", 11, INDIGO, True),
       ("connectsphere.dsl (Structurizr, C1+C2)  ·  docs/design/diagrams/*.svg (L4)  — render C1/C2 to PNG before pasting", 11, SLATE)]])

# ============================================================ 7. SPRINT CONDUCT
s = content_slide(7, "Item 4 · Process", "How the sprint was conducted")
bullets(s, Inches(0.7), Inches(1.9), Inches(7.4), Inches(5), [
    ("Two-week sprint; four Scrum ceremonies run each cycle:", 0, INK, True),
    ("Sprint Planning — pull from top of backlog into Sprint Backlog; agree a goal.", 1),
    ("Daily Scrum — short sync: done / doing / blockers.", 1),
    ("Sprint Review — demo completed stories against ACs.", 1),
    ("Retrospective — what went well / improve for next sprint.", 1),
    ("Roles — [ PO ] sets priority · [ SM ] facilitates · dev team delivers.", 0),
    ("Work tracked on the Jira board: To Do → In Progress → In Review → Done.", 0),
    ("Each story: feature branch → PR → review → CI green → merge (DoD gated).", 0, INK, True),
], size=14.5, gap=9)
placeholder(s, Inches(8.4), Inches(1.9), Inches(4.2), Inches(3.4),
            "[ image: Jira board snapshot\nor ceremony timeline ]")
text(s, Inches(8.4), Inches(5.45), Inches(4.2), Inches(1.2),
     [[("EVIDENCE: ", 11, INDIGO, True),
       ("4_Scrum_Evidence/Sprint_1/ — planning agenda + plan; meeting records", 11, SLATE)]])

# ============================================================ 8. BLOCKERS
s = content_slide(8, "Item 5 · Blockers", "Blockers & how we resolved them")
text(s, Inches(0.7), Inches(1.8), Inches(12), Inches(0.5),
     [[("Two kinds of blocker surfaced — requirement ambiguity and technical. Example of each:", 14, SLATE)]])
# card 1
rect(s, Inches(0.7), Inches(2.5), Inches(5.9), Inches(3.4), CARD, line=LIGHT)
text(s, Inches(0.95), Inches(2.7), Inches(5.4), Inches(3.1), [
    [("Requirement ambiguity", 15, INDIGO, True)],
    [("Blocker: ", 12, INK, True), ("customer gave conflicting rules — e.g. \"no pending state\" vs a request-amendment flow.", 12, SLATE)],
    [("Resolution: ", 12, GREEN, True), ("logged as customer questions J1 / J2; kept approve/reject only; flagged as open team decision rather than assuming.", 12, SLATE)],
], space_after=8)
# card 2
rect(s, Inches(6.75), Inches(2.5), Inches(5.85), Inches(3.4), CARD, line=LIGHT)
text(s, Inches(7.0), Inches(2.7), Inches(5.35), Inches(3.1), [
    [("Technical", 15, INDIGO, True)],
    [("Blocker: ", 12, INK, True), ("[ fill in: e.g. auth not ready → stories depended on login ]", 12, AMBER)],
    [("Resolution: ", 12, GREEN, True), ("[ fill in: seeded identities + role-switcher auth stub (SCRUM-13) so dev could proceed in parallel ]", 12, AMBER)],
], space_after=8)

# ============================================================ 9. CI TESTS
s = content_slide(9, "Item 6 · CI", "Test cases in the CI pipeline")
y = Inches(2.0); w = Inches(2.78); h = Inches(1.7); gap = Inches(0.18); x0 = Inches(0.7)
metric_card(s, x0, y, w, h, "203", "API unit/integration (pytest)")
metric_card(s, x0+(w+gap), y, w, h, "16", "Frontend unit (Vitest)")
metric_card(s, x0+2*(w+gap), y, w, h, "53", "E2E browser (Playwright)")
metric_card(s, x0+3*(w+gap), y, w, h, "272", "total automated", GREEN)
bullets(s, Inches(0.7), Inches(4.1), Inches(12), Inches(2.6), [
    ("GitHub Actions runs on every push & pull request — two jobs (frontend + api).", 0, INK, True),
    ("frontend job: npm ci → lint → security audit → build → Vitest unit → Playwright E2E.", 0),
    ("api job: install → ruff → spin up disposable Postgres → pytest with coverage.", 0),
    ("Merge is blocked unless CI is green — test evidence uploaded as build artifacts.", 0, INK, True),
], size=14.5, gap=9)
text(s, Inches(0.7), Inches(6.75), Inches(12), Inches(0.4),
     [[("LIVE: ", 11, INDIGO, True), ("GitHub → Actions tab (green runs)  ·  CONFIG: .github/workflows/ci.yml", 11, SLATE)]])

# ============================================================ 10. DoD
s = content_slide(10, "Item 7 · DoD", "Definition of Done — & where it lives")
bullets(s, Inches(0.7), Inches(1.9), Inches(7.3), Inches(5), [
    ("Linked to a Jira issue (SCRUM-xx).", 0),
    ("All acceptance criteria for the issue are met.", 0),
    ("Tests written and passing (unit and/or E2E).", 0),
    ("No regressions introduced.", 0),
    ("RBAC, input validation, error handling considered; no secrets committed.", 0),
    ("Documentation updated where relevant.", 0),
    ("A human reviewed & understood the code (incl. AI-generated parts).", 0),
    ("CI is green.", 0, INK, True),
], size=15, gap=10)
rect(s, Inches(8.4), Inches(1.9), Inches(4.2), Inches(4.4), CARD, line=LIGHT)
text(s, Inches(8.65), Inches(2.15), Inches(3.7), Inches(4), [
    [("Where it's recorded", 15, INDIGO, True)],
    [("CONTRIBUTING.md", 13, INK, True)],
    [("the canonical DoD text", 11, SLATE)],
    [("", 6, SLATE)],
    [(".github/pull_request_template.md", 13, INK, True)],
    [("DoD as a checklist on every PR", 11, SLATE)],
    [("", 6, SLATE)],
    [("Confluence", 13, INK, True)],
    [("team-facing ways-of-working", 11, SLATE)],
], space_after=6)

# ============================================================ 11. UNIT TESTING
s = content_slide(11, "Item 8 · Testing", "Unit testing — location & automated vs manual")
bullets(s, Inches(0.7), Inches(1.9), Inches(7.3), Inches(5), [
    ("Where it lives in the repo:", 0, INK, True),
    ("api/tests/ — pytest unit + integration (15 files, 203 tests).", 1),
    ("frontend/src/lib/*.test.ts — Vitest unit tests (pure logic).", 1),
    ("frontend/e2e/*.spec.ts — Playwright end-to-end (8 files, 53 tests).", 1),
    ("docs/SCRUM-xx.md — per-story test instructions.", 1),
    ("Automated (in CI): all pytest, Vitest, and Playwright suites — every push/PR.", 0, GREEN, True),
    ("Manual: exploratory UI checks + the PR reviewer running the story per docs/ before approving.", 0, AMBER, True),
], size=14.5, gap=9)
rect(s, Inches(8.4), Inches(1.9), Inches(4.2), Inches(4.0), CARD, line=LIGHT)
text(s, Inches(8.65), Inches(2.15), Inches(3.7), Inches(3.6), [
    [("Test pyramid", 15, INDIGO, True)],
    [("203  unit / integration (API)", 12, SLATE)],
    [(" 16  unit (frontend logic)", 12, SLATE)],
    [(" 53  end-to-end (browser)", 12, SLATE)],
    [("———————————————", 11, MUTED)],
    [("272  automated, gated by CI", 13, GREEN, True)],
    [("", 8, SLATE)],
    [("+ manual review per story", 11, AMBER)],
], space_after=7)

# ============================================================ 12. CLOSING
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, INK)
rect(s, 0, Inches(3.4), Inches(2.0), Inches(0.06), INDIGO)
text(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.2),
     [[("Where we'd value your feedback", 36, WHITE, True)]])
bullets_items = [
    ("Is our backlog altitude (DEEP) pitched right for the next sprint?", 0, RGBColor(0xC7,0xD2,0xFE)),
    ("Are the open customer conflicts (J1 / J2) handled the way you'd expect?", 0, RGBColor(0xC7,0xD2,0xFE)),
    ("[ add your own questions for the instructors here ]", 0, MUTED),
]
bullets(s, Inches(0.9), Inches(3.8), Inches(11), Inches(2.5), bullets_items, size=17, gap=14)
text(s, Inches(0.9), Inches(6.7), Inches(11), Inches(0.5),
     [[("Thank you — ConnectSphere · Team G6T4", 13, MUTED)]])

out = "/Users/marcus/Desktop/ConnectSphere/5_Slide_Deck/Sprint1_Consultation.pptx"
prs.save(out)
print("saved", out, "·", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
